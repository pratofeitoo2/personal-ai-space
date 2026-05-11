#!/usr/bin/env python3
"""
File format converters — convert .txt, .pdf, .docx to .md before ingestion.

Each converter returns (markdown_content: str, metadata: dict).
Supported: .txt, .pdf, .docx
"""
from __future__ import annotations

import subprocess
from pathlib import Path
from datetime import datetime


# ── .txt ──────────────────────────────────────────────────────────────────────

def convert_txt(file_path: Path) -> tuple[str, dict]:
    """Plain text → Markdown. Wraps content in a header using the filename."""
    text = file_path.read_text(encoding="utf-8", errors="replace")
    title = file_path.stem.replace("_", " ").replace("-", " ").strip()
    metadata = {
        "source": "txt",
        "converted_at": datetime.now().isoformat(),
        "original_name": file_path.name,
    }
    content = f"# {title}\n\n{text}"
    return content, metadata


# ── .pdf ──────────────────────────────────────────────────────────────────────

def convert_pdf(file_path: Path) -> tuple[str, dict]:
    """PDF → Markdown via pdfplumber (fallback: pypdf2 → text)."""
    import sys as _sys
    import io as _io

    metadata: dict = {
        "source": "pdf",
        "converted_at": datetime.now().isoformat(),
        "original_name": file_path.name,
    }

    # Try pdfplumber first (better text extraction)
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            pages: list[str] = []
            metadata["page_count"] = len(pdf.pages)
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(f"## Page {i + 1}\n\n{text}")
            content = "\n\n---\n\n".join(pages)
            metadata["extractor"] = "pdfplumber"
            metadata["page_count"] = len(pdf.pages)
    except ImportError:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            pages = []
            metadata["page_count"] = len(reader.pages)
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(f"## Page {i + 1}\n\n{text}")
            content = "\n\n---\n\n".join(pages)
            metadata["extractor"] = "PyPDF2"
        except ImportError:
            raise RuntimeError(
                "Neither pdfplumber nor PyPDF2 is installed. "
                "Run: pip install pdfplumber  # or pip install PyPDF2"
            )

    title = file_path.stem.replace("_", " ").replace("-", " ").strip()
    return f"# {title}\n\n{content}", metadata


# ── .docx ─────────────────────────────────────────────────────────────────────

def convert_docx(file_path: Path) -> tuple[str, dict]:
    """DOCX → Markdown via python-docx."""
    from docx import Document

    doc = Document(str(file_path))
    metadata: dict = {
        "source": "docx",
        "converted_at": datetime.now().isoformat(),
        "original_name": file_path.name,
        "paragraph_count": len(doc.paragraphs),
    }

    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue
        style_name = (para.style.name or "").lower()
        if "heading" in style_name or "title" in style_name:
            level = 1
            for lvl in range(1, 10):
                if str(lvl) in style_name:
                    level = lvl
                    break
            lines.append(f"{'#' * level} {text}")
        elif para.runs and para.runs[0].bold:
            lines.append(f"**{text}**")
        else:
            lines.append(text)

    # Extract tables
    tables_md: list[str] = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append("| " + " | ".join(cells) + " |")
        if rows:
            # Header separator
            rows.insert(1, "| " + " | ".join(["---"] * len(table.columns)) + " |")
            tables_md.append("\n".join(rows) + "\n")

    content = "\n".join(lines)
    if tables_md:
        content += "\n\n## Tables\n\n" + "\n\n".join(tables_md)

    return f"# {file_path.stem.replace('_', ' ').replace('-', ' ').strip()}\n\n{content}", metadata


# ── Router ────────────────────────────────────────────────────────────────────

CONVERTERS = {
    ".txt": convert_txt,
    ".pdf": convert_pdf,
    ".docx": convert_docx,
}


def convert_to_markdown(file_path: Path) -> tuple[Path, dict]:
    """
    Convert a non-markdown file to .md and write it next to the original.

    Returns:
        (path_to_md_file, conversion_metadata)

    Raises:
        ValueError if extension is unsupported.
    """
    ext = file_path.suffix.lower()
    if ext not in CONVERTERS:
        raise ValueError(f"Unsupported conversion extension: {ext}")

    content, metadata = CONVERTERS[ext](file_path)

    # Write .md next to original with same stem + .md
    md_path = file_path.with_suffix(".md")
    md_path.write_text(content, encoding="utf-8")
    return md_path, metadata